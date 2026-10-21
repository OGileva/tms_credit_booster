from __future__ import annotations

import csv
import io
import json
import shutil
import tempfile
import unittest
from contextlib import redirect_stderr, redirect_stdout
from dataclasses import replace
from pathlib import Path

from docx import Document
from openpyxl import Workbook, load_workbook

from qa_cases.cli import main
from qa_cases.core import (
    CANONICAL_HEADERS,
    METADATA_HEADERS,
    TestCase,
    create_section,
    export_cases,
    import_cases,
    load_sections,
    render_table,
    resolve_section,
    resolve_section_queries,
    resolve_tags,
    select_cases,
    validate_repository,
    write_case,
)


class RepositoryTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        shutil.copytree(Path(__file__).resolve().parents[1] / "schemas", self.root / "schemas")
        (self.root / "project.json").write_text("{}", encoding="utf-8")
        section = self.root / "cases" / "test-product" / "demo" / "sample-section"
        section.mkdir(parents=True)
        (section / "section.json").write_text(
            json.dumps(
                {
                    "name": "Тестовый раздел",
                    "aliases": ["Пример раздела"],
                    "id_prefix": "TEST-DEMO-CASE",
                },
                ensure_ascii=False,
            ),
            encoding="utf-8",
        )
        dictionaries = self.root / "dictionaries"
        dictionaries.mkdir()
        (dictionaries / "tags.json").write_text(
            json.dumps(
                {
                    "tags": [
                        {"id": "regression", "aliases": ["регресс"], "description": "Регрессионные проверки"},
                        {"id": "positive", "aliases": ["позитивный"], "description": "Позитивные проверки"},
                        {"id": "critical", "aliases": ["критичный"], "description": "Критичные проверки"},
                    ]
                },
                ensure_ascii=False,
            ),
            encoding="utf-8",
        )

    def tearDown(self):
        self.temp.cleanup()

    def _case(self, case_id: str = "1", title: str = "Проверка тестового действия") -> TestCase:
        return TestCase(
            id=case_id,
            title=title,
            description="Проверка обработки тестовых данных",
            preconditions="1. Подготовлены тестовые данные",
            steps="1. Выполнить тестовое действие",
            expected_result="1. Получен тестовый результат",
            status="active",
            tags=("regression",),
            tasks=("CBN-1",),
            source="cases.csv",
            product="test-product",
            section_path="demo/sample-section",
            section_name="Тестовый раздел",
            file=f"cases/test-product/demo/sample-section/{case_id}.md",
        )

    def test_markdown_case_is_valid(self):
        case = self._case()
        write_case(self.root / case.file, case)
        cases, issues = validate_repository(self.root)
        self.assertEqual([], issues)
        self.assertEqual("1", cases[0].id)
        self.assertEqual("Тестовый раздел", cases[0].section_name)
        self.assertEqual(case, cases[0])

    def test_mass_csv_import_allocates_ids_and_resolves_tag_alias(self):
        csv_path = self.root / "input.csv"
        with csv_path.open("w", encoding="utf-8", newline="") as handle:
            writer = csv.writer(handle)
            writer.writerow(["Name", "Description", "Preconditions", "Steps", "Expected Result", "Tags"])
            writer.writerow(["Первый кейс", "Описание 1", "Клиент создан", "Открыть раздел", "Раздел открыт", "регресс"])
            writer.writerow(["Второй кейс", "Описание 2", "Подготовлены данные", "Обработать данные", "Данные обработаны", "позитивный"])
        section = resolve_section(self.root, "Тестовый раздел", "test-product")
        imported = import_cases(
            self.root,
            source_path=csv_path,
            section=section,
            command_tags=("regression",),
            tasks=("CBN-2",),
        )
        self.assertEqual(["1", "2"], [case.id for case in imported])
        self.assertTrue((section.directory / "2.md").exists())
        cases, issues = validate_repository(self.root)
        self.assertEqual([], issues)
        self.assertEqual(2, len(cases))
        self.assertEqual(("positive", "regression"), cases[1].tags)

    def test_selection_by_russian_section_and_tag(self):
        case = self._case()
        selected = select_cases(
            [case],
            sections=("Тестовый раздел",),
            tags=resolve_tags(self.root, ["регресс"]),
        )
        self.assertEqual([case], selected)
        table = render_table(selected)
        self.assertIn("1", table)
        self.assertIn("Expected Result", table)

    def test_xlsx_export_keeps_five_canonical_columns(self):
        output = self.root / "exports" / "selection.xlsx"
        export_cases(output, [self._case()], "xlsx")
        workbook = load_workbook(output, read_only=False, data_only=False)
        worksheet = workbook["Тест-кейсы"]
        self.assertEqual(5, worksheet.max_column)
        self.assertEqual(
            ["Name", "Description", "Preconditions", "Steps", "Expected Result"],
            [worksheet.cell(1, column).value for column in range(1, 6)],
        )
        self.assertEqual("1 — Проверка тестового действия", worksheet["A2"].value)
        workbook.close()

    def test_duplicate_id_is_rejected(self):
        first = self._case()
        second = self._case(title="Другое название")
        write_case(self.root / first.file, first)
        duplicate_path = self.root / "cases" / "test-product" / "demo" / "sample-section" / "copy.md"
        write_case(duplicate_path, second)
        _, issues = validate_repository(self.root)
        self.assertTrue(any("имя файла" in issue.message or "ID дублируется" in issue.message for issue in issues))

    def _input_record(self, name="Новый кейс", **values):
        return dict(zip(CANONICAL_HEADERS, [name, "Описание", "Создан клиент", "Открыть раздел", "Раздел открыт"]), **values)

    def _import_records(self, records, dry_run=False):
        source = self.root / "input.json"
        source.write_text(json.dumps(records, ensure_ascii=False), encoding="utf-8")
        return import_cases(
            self.root, source_path=source, section=resolve_section(self.root, "Тестовый раздел"), dry_run=dry_run
        )

    def _cli(self, *args):
        stdout, stderr = io.StringIO(), io.StringIO()
        with redirect_stdout(stdout), redirect_stderr(stderr):
            code = main(["--root", str(self.root), *args])
        return code, stdout.getvalue(), stderr.getvalue()

    def test_metadata_schema_rejects_missing_unknown_and_mistyped_fields(self):
        case = self._case()
        path = self.root / case.file
        write_case(path, case)
        original = path.read_text(encoding="utf-8")
        variants = [
            ('tasks = ["CBN-1"]\n', "", "tasks"),
            ('tags = ["regression"]\n', "", "tags"),
            ('tasks = ["CBN-1"]', 'tasks = "CBN-1"', "tasks"),
            ('tasks = ["CBN-1"]', 'tasks = [1]', "tasks"),
            ('tasks = ["CBN-1"]', 'tasks = ["CBN-1", "CBN-1"]', "tasks"),
            ('title = "Проверка тестового действия"', 'title = 123', "title"),
            ('source = "cases.csv"', 'source = 123', "source"),
            ('source = "cases.csv"', 'unexpected = "value"', "unexpected"),
        ]
        for before, after, field in variants:
            with self.subTest(field=field, after=after):
                path.write_text(original.replace(before, after), encoding="utf-8")
                cases, issues = validate_repository(self.root)
                self.assertEqual([], cases)
                self.assertTrue(any(field in issue.message for issue in issues), issues)

    def test_create_section_rejects_duplicate_prefix_and_empty_name_without_writing(self):
        for name, prefix in [("Другой раздел", "TEST-DEMO-CASE"), ("  ", "TEST-NEW")]:
            with self.subTest(name=name):
                with self.assertRaises(ValueError):
                    create_section(self.root, product="test-product", path="other", name=name, id_prefix=prefix)
                self.assertFalse((self.root / "cases/test-product/other").exists())

    def test_create_section_and_validate_reject_invalid_config(self):
        section = create_section(self.root, product="test-product", path="other", name="Другой раздел", id_prefix="TEST-NEW")
        _, issues = validate_repository(self.root)
        self.assertEqual([], issues)
        config = section.directory / "section.json"
        for payload in [
            {"name": "Другой раздел", "id_prefix": "TEST-DEMO-CASE"},
            {"name": "", "id_prefix": "TEST-NEW"},
            {"name": "Другой раздел", "id_prefix": "TEST-NEW", "aliases": "alias"},
        ]:
            with self.subTest(payload=payload):
                config.write_text(json.dumps(payload), encoding="utf-8")
                _, issues = validate_repository(self.root)
                self.assertTrue(issues)

    def test_tag_dictionary_rejects_duplicate_ids_and_ambiguous_aliases(self):
        path = self.root / "dictionaries/tags.json"
        for tags in [
            [{"id": "regression", "description": "Регресс"}] * 2,
            [{"id": "regression", "description": "Регресс", "aliases": ["positive"]},
             {"id": "positive", "description": "Позитивный"}],
            [{"id": "regression", "description": "Регресс", "aliases": "регресс"}],
        ]:
            with self.subTest(tags=tags):
                path.write_text(json.dumps({"tags": tags}), encoding="utf-8")
                _, issues = validate_repository(self.root)
                self.assertTrue(any(issue.source == "dictionaries/tags.json" for issue in issues))

    def test_cli_selects_regression_critical_and_their_intersection(self):
        cases = [
            self._case(),
            replace(self._case("2", "Критичный регресс"), tags=("critical", "regression")),
            replace(self._case("3", "Критичный вне регресса"), tags=("critical",)),
            replace(self._case("4", "Черновик"), tags=("critical", "regression"), status="draft"),
            replace(self._case("5", "Устаревший"), tags=("critical", "regression"), status="deprecated"),
        ]
        for case in cases:
            write_case(self.root / case.file, case)
        filters = [
            (("--tag", "регресс"), [cases[0].id, cases[1].id]),
            (("--tag", "критичный"), [cases[1].id, cases[2].id]),
            (("--tag", "регресс", "--tag", "критичный"), [cases[1].id]),
            (("--tag", "критичный", "--status", "all"), [case.id for case in cases[1:]]),
        ]
        for args, expected_ids in filters:
            with self.subTest(filters=args):
                code, output, error = self._cli("list", *args)
                self.assertEqual(0, code, error)
                self.assertEqual(expected_ids, [line.split("\t")[0] for line in output.splitlines()[1:]])
        code, output, error = self._cli("export", "--tag", "регресс", "--tag", "критичный", "--include-metadata")
        self.assertEqual(0, code, error)
        self.assertIn("| ID | Tags | Status | Tasks |", output)
        self.assertIn(cases[1].id, output)
        data_rows = output.splitlines()[2:]
        self.assertEqual(1, len(data_rows))
        self.assertIn(f"| {cases[1].id} |", data_rows[0])

    def test_docs_do_not_enter_case_validation_or_export(self):
        document = self.root / "docs" / "test-product" / "requirements.md"
        document.parent.mkdir(parents=True)
        document.write_text("# Требования\n\nПроизвольный Markdown-документ.", encoding="utf-8")
        case = self._case()
        write_case(self.root / case.file, case)
        self.assertEqual(([case], []), validate_repository(self.root))
        code, output, error = self._cli("export")
        self.assertEqual(0, code, error)
        self.assertIn(case.id, output)
        self.assertNotIn("Произвольный Markdown-документ", output)

    def test_metadata_exports_preserve_id_tags_status_and_tasks_on_import(self):
        case = replace(self._case(), tags=("critical", "regression"), status="draft", tasks=("CBN-1", "CBN-2"))
        for format_name, suffix in [("xlsx", "xlsx"), ("csv", "csv"), ("markdown", "md"), ("json", "json")]:
            with self.subTest(format=format_name):
                path = self.root / f"metadata.{suffix}"
                export_cases(path, [case], format_name, include_metadata=True)
                if format_name == "xlsx":
                    workbook = load_workbook(path)
                    worksheet = workbook.active
                    self.assertEqual(CANONICAL_HEADERS + METADATA_HEADERS, [cell.value for cell in worksheet[1]])
                    self.assertEqual("A1:J2", worksheet.tables["TestCasesTable"].ref)
                    self.assertTrue(worksheet["I2"].alignment.wrap_text)
                    workbook.close()
                imported = import_cases(
                    self.root, source_path=path, section=resolve_section(self.root, "Тестовый раздел"), dry_run=True
                )
                self.assertEqual(1, len(imported))
                self.assertEqual(case.table_row(include_metadata=True), imported[0].table_row(include_metadata=True))

    def test_cli_list_and_export_resolve_section_alias(self):
        case = self._case()
        write_case(self.root / case.file, case)
        for command in ("list", "export"):
            code, output, error = self._cli(command, "--section", "Пример раздела", "--tag", "регресс")
            self.assertEqual(0, code, error)
            self.assertIn(case.id, output)

    def test_section_alias_ambiguity_requires_product_and_preserves_globs(self):
        other = create_section(self.root, product="other-product", path="demo/sample-section", name="Тестовый раздел",
                               id_prefix="OTHER-DEMO", aliases=["Пример раздела"])
        with self.assertRaisesRegex(ValueError, "неоднозначен"):
            resolve_section_queries(self.root, ["Пример раздела"])
        resolved = resolve_section_queries(self.root, ["Пример раздела"], ["test-product"])
        self.assertEqual([self._case()], select_cases([self._case()], sections=resolved))
        self.assertEqual(("demo/*",), resolve_section_queries(self.root, ["demo/*"]))
        self.assertEqual(2, self._cli("export", "--section", "Пример раздела")[0])
        self.assertEqual(other.path, "demo/sample-section")

    def test_unknown_section_returns_actionable_error(self):
        code, _, error = self._cli("export", "--section", "Неизвестный раздел")
        self.assertEqual(2, code)
        self.assertIn("qa-cases sections", error)

    def test_import_ignores_input_ids_and_assigns_sequential_ids(self):
        records = [self._input_record("Первый", ID="999"), self._input_record("Второй", ID="1")]
        imported = self._import_records(records)
        self.assertEqual(["1", "2"], [case.id for case in imported])

    def test_import_skips_duplicates_by_required_field_content(self):
        first = self._input_record("Повтор", ID="999")
        imported = self._import_records([first])
        duplicate = self._input_record("Повтор", ID="1")
        duplicate["Tags"] = "critical"
        duplicate["Description"] = "Другое описание"
        result = self._import_records([duplicate, duplicate, self._input_record("Новый")])
        self.assertEqual(["2"], [case.id for case in result])
        self.assertEqual("Новый", result[0].title)

    def test_deleted_ids_are_never_reused(self):
        write_case(self.root / self._case("1").file, self._case("1"))
        write_case(self.root / self._case("3").file, self._case("3", "Третий"))
        imported = self._import_records([self._input_record("Новый")])
        self.assertEqual("4", imported[0].id)

    def test_id_range_exhaustion_writes_no_pending_cases(self):
        case = self._case("999998")
        write_case(self.root / case.file, case)
        for dry_run in (True, False):
            with self.subTest(dry_run=dry_run):
                with self.assertRaisesRegex(ValueError, "исчерпан"):
                    self._import_records([self._input_record("Первый"), self._input_record("Второй")], dry_run=dry_run)
                self.assertEqual([case.id], [case.id for case in validate_repository(self.root)[0]])
        result = self._import_records([self._input_record("Последний")])
        self.assertEqual("999999", result[0].id)
        with self.assertRaisesRegex(ValueError, "исчерпан"):
            self._import_records([self._input_record("Следующий")])
        self.assertEqual([], validate_repository(self.root)[1])

    def test_missing_required_field_reports_row_and_writes_nothing(self):
        for field in ("Name", "Preconditions", "Steps", "Expected Result"):
            with self.subTest(field=field):
                invalid = self._input_record("Неполный")
                invalid[field] = ""
                with self.assertRaisesRegex(ValueError, "input.json:строка 2"):
                    self._import_records([self._input_record("Полный"), invalid])
                self.assertEqual([], list((self.root / "cases").rglob("*.md")))

    def test_description_is_optional_on_import(self):
        record = self._input_record("Без описания")
        record["Description"] = ""
        imported = self._import_records([record], dry_run=True)
        self.assertEqual("", imported[0].description)

    def test_xlsx_missing_column_on_one_sheet_rejects_entire_import(self):
        path = self.root / "input.xlsx"
        workbook = Workbook()
        for worksheet, headers in [(workbook.active, CANONICAL_HEADERS), (workbook.create_sheet("Неполный"), CANONICAL_HEADERS[:-1])]:
            worksheet.append(headers)
            worksheet.append(list(self._input_record().values())[:len(headers)])
        workbook.save(path)
        with self.assertRaisesRegex(ValueError, "input.xlsx:Неполный"):
            import_cases(self.root, source_path=path, section=resolve_section(self.root, "Тестовый раздел"))
        self.assertEqual([], list((self.root / "cases").rglob("*.md")))

    def test_docx_missing_column_in_one_table_rejects_entire_import(self):
        path = self.root / "input.docx"
        document = Document()
        for headers in (CANONICAL_HEADERS, CANONICAL_HEADERS[:-1]):
            table = document.add_table(rows=1, cols=len(headers))
            for cell, value in zip(table.rows[0].cells, headers):
                cell.text = value
            for cell, value in zip(table.add_row().cells, self._input_record().values()):
                cell.text = value
        document.save(path)
        with self.assertRaisesRegex(ValueError, "input.docx:таблица-2"):
            import_cases(self.root, source_path=path, section=resolve_section(self.root, "Тестовый раздел"))
        self.assertEqual([], list((self.root / "cases").rglob("*.md")))

    def test_xlsx_empty_leading_rows_do_not_hide_cases(self):
        path = self.root / "input.xlsx"
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.append([None] * 5)
        worksheet.append(CANONICAL_HEADERS)
        worksheet.append(list(self._input_record().values()))
        workbook.create_sheet("Пустой")
        workbook.save(path)
        cases = import_cases(self.root, source_path=path, section=resolve_section(self.root, "Тестовый раздел"))
        self.assertEqual(1, len(cases))

    def test_xlsx_imports_cases_from_all_valid_sheets(self):
        path = self.root / "input.xlsx"
        workbook = Workbook()
        for index, worksheet in enumerate((workbook.active, workbook.create_sheet("Второй")), start=1):
            worksheet.append(CANONICAL_HEADERS)
            worksheet.append(list(self._input_record(f"Кейс {index}").values()))
        workbook.save(path)
        result = import_cases(self.root, source_path=path, section=resolve_section(self.root, "Тестовый раздел"))
        self.assertEqual(["Кейс 1", "Кейс 2"], [case.title for case in result])
        self.assertEqual([], validate_repository(self.root)[1])

    def test_docx_imports_cases_from_all_valid_tables(self):
        path = self.root / "input.docx"
        document = Document()
        for index in (1, 2):
            table = document.add_table(rows=1, cols=5)
            for cell, value in zip(table.rows[0].cells, CANONICAL_HEADERS):
                cell.text = value
            for cell, value in zip(table.add_row().cells, self._input_record(f"Кейс {index}").values()):
                cell.text = value
        document.save(path)
        result = import_cases(self.root, source_path=path, section=resolve_section(self.root, "Тестовый раздел"))
        self.assertEqual(["Кейс 1", "Кейс 2"], [case.title for case in result])
        self.assertEqual([], validate_repository(self.root)[1])

    def test_tsv_import_preserves_tabs_and_newlines_in_quoted_fields(self):
        path = self.root / "input.tsv"
        case = replace(self._case(), description="Значения\tA | B\nВторая строка")
        with path.open("w", encoding="utf-8", newline="") as handle:
            writer = csv.writer(handle, delimiter="\t")
            writer.writerow(CANONICAL_HEADERS)
            writer.writerow(case.table_row())
        result = import_cases(self.root, source_path=path, section=resolve_section(self.root, "Тестовый раздел"))
        self.assertEqual(case.table_row(), result[0].table_row())

    def test_xlsx_incomplete_row_reports_sheet_and_actual_row(self):
        path = self.root / "input.xlsx"
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Кейсы"
        worksheet.append(CANONICAL_HEADERS)
        worksheet.append(list(self._input_record("Первый").values()))
        worksheet.append(list(self._input_record("Неполный").values())[:-1])
        workbook.save(path)
        with self.assertRaisesRegex(ValueError, "input.xlsx:Кейсы:строка 3.*Expected Result"):
            import_cases(self.root, source_path=path, section=resolve_section(self.root, "Тестовый раздел"))
        self.assertEqual([], list((self.root / "cases").rglob("*.md")))

    def test_export_import_roundtrip_preserves_all_five_fields(self):
        case = replace(self._case(), description="Значения A | B, путь C:\\files\\, <br> и &lt;br&gt;\nВторая строка")
        for format_name, suffix in [("xlsx", "xlsx"), ("csv", "csv"), ("json", "json"), ("markdown", "md")]:
            with self.subTest(format=format_name):
                path = self.root / f"roundtrip.{suffix}"
                export_cases(path, [case], format_name)
                result = import_cases(self.root, source_path=path, section=resolve_section(self.root, "Тестовый раздел"), dry_run=True)
                self.assertEqual([case.table_row()], [item.table_row() for item in result])

    def test_markdown_escaped_pipes_backslashes_and_edge_pipes_preserve_content(self):
        for value in ["|A|", "A \\| B", "A \\\\| B", "C:\\path\\", "<br /> & <b>текст</b>"]:
            with self.subTest(value=value):
                case = replace(self._case(), description=value)
                path = self.root / "roundtrip.md"
                export_cases(path, [case], "markdown")
                result = import_cases(self.root, source_path=path, section=resolve_section(self.root, "Тестовый раздел"), dry_run=True)
                self.assertEqual(case.table_row(), result[0].table_row())

    def test_xlsx_formula_like_description_remains_literal_text(self):
        case = replace(self._case(), description="=1+1")
        path = self.root / "literal.xlsx"
        export_cases(path, [case], "xlsx")
        for data_only in (True, False):
            workbook = load_workbook(path, data_only=data_only)
            self.assertEqual("=1+1", workbook.active["B2"].value)
            self.assertEqual("s", workbook.active["B2"].data_type)
            workbook.close()
        result = import_cases(self.root, source_path=path, section=resolve_section(self.root, "Тестовый раздел"))
        self.assertEqual(case.table_row(), result[0].table_row())

    def test_cli_import_dry_run_then_export_without_modifying_cases(self):
        source = self.root / "input.json"
        source.write_text(json.dumps([self._input_record()]), encoding="utf-8")
        args = ["import", "--source", str(source), "--section", "Пример раздела", "--tag", "регресс", "--task", "CBN-42"]
        code, output, error = self._cli(*args, "--dry-run")
        self.assertEqual(0, code, error)
        self.assertIn("Проверено без записи: 1", output)
        self.assertEqual([], list((self.root / "cases").rglob("*.md")))
        self.assertEqual(0, self._cli(*args)[0])
        snapshot = {path: path.read_bytes() for path in (self.root / "cases").rglob("*") if path.is_file()}
        for format_name, suffix in [("xlsx", "xlsx"), ("csv", "csv"), ("json", "json"), ("markdown", "md")]:
            with self.subTest(format=format_name):
                code, _, error = self._cli("export", "--section", "Пример раздела", "--format", format_name,
                                           "--output", f"exports/selected.{suffix}")
                self.assertEqual(0, code, error)
                self.assertTrue((self.root / f"exports/selected.{suffix}").exists())
        self.assertEqual(snapshot, {path: path.read_bytes() for path in (self.root / "cases").rglob("*") if path.is_file()})


if __name__ == "__main__":
    unittest.main()
